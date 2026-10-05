#!/usr/bin/env python3
"""Build the matching package for the terminal: numbered floor plans (PDF) + a matching workbook (XLSX).

For every floor plan the CAD rooms are numbered; the user fills in, for every storage location
of the asset register, which numbered room it is.

Output (sandbox/indoor-map/mapping/, gitignored - contains asset counts):
  航廈對照圖_<floor>.pdf   overview page + detail pages, vector (zoom freely), grid + north arrow
  航廈配對表.xlsx          sheet 配對表 (fill in) + one candidate list per floor + 說明

Orientation: the drawing's LEFT edge is north (north arrow points left).
"""
import json
import math
import re
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.collections import PatchCollection
from matplotlib.patches import PathPatch, Polygon as MPoly
from matplotlib.path import Path as MPath
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from shapely.geometry import Point, Polygon

sys.stdout.reconfigure(encoding="utf-8")
plt.rcParams["font.sans-serif"] = ["Microsoft JhengHei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["pdf.fonttype"] = 42

BASE = Path(__file__).resolve().parents[1]
DATA = BASE / "data"
OUT = BASE / "mapping"
FLOORS = [("P", "B1 停車場", "terminal-b2", "term-b2"), ("S", "淺挖區", "terminal-b1", "term-b1"),
          ("1F", "1F", "terminal-1f", "term-1f"), ("2F", "2F", "terminal-2f", "term-2f"),
          ("3F", "3F", "terminal-3f", "term-3f")]
NOISE = re.compile(r"^(TYPE|金屬|落水管|排水|水溝|採光井|蓋板|地板|RC|\d+\s*cm|20cm|平台|伸縮|電動|防火|鐵捲門|等\s*待|等侯|"
                   r"圓鋼|鋼$|柱$|實際尺寸|鉸鋅|輸送|手推車|消防箱|植栽|景觀|集水井|GRP|40E|30E|6\"|50 PVC|R\d+|[A-Z]\d+$|"
                   r"[甲乙丙丁]$|.*樓梯$|.*電梯$|.*電扶梯|.*詳$|.*詳圖$|中山室|\.電扶梯|接梯門|活動空橋|逃生|緊急|落水|"
                   r"飲料|自動販|自動提|提款|行動電話|公共電話|飲水機|化妝台|等候線)")
GRID = 2000  # cm


def load(plan):
    return json.loads((DATA / f"{plan}.json").read_text(encoding="utf-8"))


def wall_path(plan):
    svg = (DATA / f"{plan}.svg").read_text(encoding="utf-8")
    m = re.search(r'<g id="wall"[^>]*><path[^>]*d="([^"]*)"', svg)
    verts, codes = [], []
    for sub in m.group(1).split("M")[1:]:
        try:
            pts = [tuple(float(v) for v in p.replace("Z", "").split()) for p in sub.strip().split(" L")]
        except ValueError:
            continue
        pts = [(x, -y) for x, y in pts]
        if len(pts) < 3:
            continue
        verts += pts + [pts[0]]
        codes += [MPath.MOVETO] + [MPath.LINETO] * (len(pts) - 1) + [MPath.CLOSEPOLY]
    return MPath(verts, codes)


def candidates(prefix, meta):
    """Numbered candidates: auto-detected rooms + room-like labels outside every room."""
    faces = []
    for r in meta.get("rooms", []):
        ring = r["polygon"][0]
        poly = Polygon(ring)
        if not poly.is_valid or not (2e4 <= poly.area <= 1.5e7):
            continue
        faces.append({"poly": poly, "name": r.get("name") or "", "number": r.get("number") or "", "area": poly.area / 1e4})
    inside = lambda pt: any(f["poly"].contains(pt) for f in faces)
    points = []
    seen = set()
    for l in meta["labels"]:
        t = l["text"].strip()
        if re.fullmatch(r"[\d\-]+", t) or NOISE.match(t) or len(t) < 2:
            continue
        pt = Point(l["x"], l["y"])
        if inside(pt):
            continue
        key = (t, round(l["x"] / 300), round(l["y"] / 300))
        if key in seen:
            continue
        seen.add(key)
        points.append({"poly": None, "pt": (l["x"], l["y"]), "name": t, "number": "", "area": 0.0})
    items = faces + points
    # numbering: along the drawing's x axis (north -> south), then y
    items.sort(key=lambda c: (round((c["poly"].centroid.x if c["poly"] else c["pt"][0]) / 600),
                              c["poly"].centroid.y if c["poly"] else c["pt"][1]))
    for i, c in enumerate(items, 1):
        c["id"] = f"{prefix}-{i:03d}"
        c["xy"] = (c["poly"].representative_point().x, c["poly"].representative_point().y) if c["poly"] else c["pt"]
    return items


def decorate(ax, bounds, scale_ref=1000):
    x0, y0, x1, y1 = bounds
    for gx in range(max(0, math.floor(x0 / GRID) * GRID), int(x1) + 1, GRID):
        ax.axvline(gx, color="#7aa7d9", lw=.35, alpha=.7, zorder=1)
        ax.text(gx + 40, y1 - 60, f"{chr(65 + int(gx // GRID) % 26)}", color="#2b6cb0", fontsize=8, va="top", zorder=6)
    for gy in range(max(0, math.floor(y0 / GRID) * GRID), int(y1) + 1, GRID):
        ax.axhline(gy, color="#7aa7d9", lw=.35, alpha=.7, zorder=1)
        ax.text(x0 + 40, gy + 40, f"{int(gy // GRID)}", color="#2b6cb0", fontsize=8, va="bottom", zorder=6)
    # north arrow (north is the LEFT edge)
    w, h = x1 - x0, y1 - y0
    ax.annotate("", xy=(x0 + .84 * w, y0 + .94 * h), xytext=(x0 + .93 * w, y0 + .94 * h),
                arrowprops=dict(arrowstyle="-|>", lw=2.2, color="#c0392b"), zorder=8)
    ax.text(x0 + .835 * w, y0 + .955 * h, "北", color="#c0392b", fontsize=13, fontweight="bold", ha="right", zorder=8)
    ax.plot([x1 - scale_ref - .02 * w, x1 - .02 * w], [y0 + .03 * h] * 2, color="k", lw=3, zorder=8)
    ax.text(x1 - scale_ref / 2 - .02 * w, y0 + .045 * h, "10 m", ha="center", fontsize=8, zorder=8)


def draw(ax, wall, cands, bounds, fs, show_name=True):
    ax.add_patch(PathPatch(wall, fc="#9aa0aa", ec="#6b7280", lw=.25, zorder=2))
    cm = plt.get_cmap("tab20")
    for i, c in enumerate(cands):
        x, y = c["xy"]
        if not (bounds[0] <= x <= bounds[2] and bounds[1] <= y <= bounds[3]):
            continue
        if c["poly"] is not None:
            ax.add_patch(MPoly(list(c["poly"].exterior.coords), closed=True, fc=cm(i % 20), alpha=.35,
                               ec="#222", lw=.7, zorder=3))
        else:
            ax.plot(x, y, marker="D", ms=5, mfc="#f59e0b", mec="#7c2d12", zorder=4)
        label = c["id"].split("-")[1] + (f" {c['name']}" if show_name and c["name"] else "")
        if show_name and c["number"]:
            label += f"({c['number']})"
        ax.text(x, y, label, fontsize=fs, ha="center", va="center", zorder=7,
                bbox=dict(boxstyle="round,pad=0.12", fc="white", ec="#444", lw=.4, alpha=.88))


def make_pdf(code, title, plan, cands, wall):
    allx = [c["xy"][0] for c in cands]
    ally = [c["xy"][1] for c in cands]
    bx = (min(allx) - 800, min(ally) - 800, max(allx) + 800, max(ally) + 800)
    path = OUT / f"航廈對照圖_{title.replace(' ', '')}.pdf"
    TW, TH = 6500, 4600
    with PdfPages(path) as pdf:
        fig, ax = plt.subplots(figsize=(16.5, 16.5 * (bx[3] - bx[1]) / (bx[2] - bx[0]) + .8))
        draw(ax, wall, cands, bx, 4.5, show_name=False)
        decorate(ax, bx)
        ax.set_xlim(bx[0], bx[2]); ax.set_ylim(bx[1], bx[3]); ax.set_aspect("equal")
        ax.set_title(f"{title} 總覽（只標編號；詳細請看後面各頁）  —  北在左邊", fontsize=12)
        pdf.savefig(fig); plt.close(fig)
        pages = 0
        y = bx[1]
        while y < bx[3]:
            x = bx[0]
            while x < bx[2]:
                tile = (x, y, x + TW, y + TH)
                if any(tile[0] <= c["xy"][0] <= tile[2] and tile[1] <= c["xy"][1] <= tile[3] for c in cands):
                    fig, ax = plt.subplots(figsize=(16.5, 16.5 * TH / TW + .6))
                    draw(ax, wall, cands, tile, 6.5)
                    decorate(ax, tile)
                    ax.set_xlim(tile[0], tile[2]); ax.set_ylim(tile[1], tile[3]); ax.set_aspect("equal")
                    ax.set_title(f"{title}  第 {pages + 1} 頁   格線 {GRID // 100} m   橘色菱形＝只有房名位置、沒有房間範圍", fontsize=10)
                    pdf.savefig(fig); plt.close(fig)
                    pages += 1
                x += TW - 700
            y += TH - 600
    print(f"{path.name}: {len(cands)} candidates, {pages} detail pages")


def main():
    OUT.mkdir(exist_ok=True)
    locs = json.loads((DATA / "locations.json").read_text(encoding="utf-8"))
    terminal = [x for x in locs if x["building"] in ("航廈", "空橋")]
    placed = json.loads((BASE / "schematic" / "terminal.json").read_text(encoding="utf-8"))
    guess = {}  # code -> (floor title, center)
    for fl in placed["floors"]:
        for sp in fl["spaces"]:
            if sp.get("rects"):
                x0, y0, x1, y1 = sp["rects"][0]
                guess[sp["code"]] = (fl["id"], ((x0 + x1) / 2, (y0 + y1) / 2))
            elif sp.get("polygons"):
                pts = sp["polygons"][0]
                guess[sp["code"]] = (fl["id"], (sum(p[0] for p in pts) / len(pts), sum(p[1] for p in pts) / len(pts)))

    wb = Workbook()
    ws = wb.active
    ws.title = "配對表"
    head = ["存置地點代碼", "存置地點名稱", "財物數", "資料推測樓層", "我先前的推測（編號）", "圖上編號（請填）", "樓層（若與左側不同請填）", "備註／位置描述"]
    ws.append(head)
    cand_by_floor = {}
    for prefix, title, plan, fid in FLOORS:
        meta = load(plan)
        cands = candidates(prefix, meta)
        cand_by_floor[fid] = (prefix, title, cands)
        make_pdf(prefix, title, plan, cands, wall_path(plan))
        sh = wb.create_sheet(f"候選_{title.replace(' ', '')}")
        sh.append(["編號", "CAD 房名", "CAD 房號", "面積(m²)", "類型", "座標 x(cm)", "座標 y(cm)"])
        for c in cands:
            sh.append([c["id"], c["name"], c["number"], round(c["area"], 1) if c["poly"] else "",
                       "房間範圍" if c["poly"] else "只有房名位置", round(c["xy"][0]), round(c["xy"][1])])
        for col, w in zip("ABCDEFG", (10, 26, 10, 10, 14, 12, 12)):
            sh.column_dimensions[col].width = w
        sh.freeze_panes = "A2"

    floor_order = {"B1": 0, "1F": 1, "2F": 2, "3F": 3, "?": 4}
    for x in sorted(terminal, key=lambda x: (floor_order.get(x["floor"], 5), x["code"])):
        g = guess.get(x["code"])
        gid = ""
        if g:
            fid, (cx, cy) = g
            prefix, title, cands = cand_by_floor[fid]
            hit = [c for c in cands if c["poly"] is not None and c["poly"].contains(Point(cx, cy))]
            if hit:
                gid = min(hit, key=lambda c: c["poly"].area)["id"]
            else:
                near = min(cands, key=lambda c: math.dist(c["xy"], (cx, cy)))
                gid = near["id"] + "（附近）"
        ws.append([x["code"].lstrip("*"), x["name"], x["count"], x["floor"], gid, "", "", ""])
    fill = PatternFill("solid", fgColor="FFF3CD")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E2F3")
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=8):
        for c in row:
            c.fill = fill
    for col, w in zip("ABCDEFGH", (12, 34, 8, 12, 18, 18, 20, 36)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    note = wb.create_sheet("說明", 0)
    for line in [
        "航廈存置地點 ↔ 圖面位置 對照表",
        "",
        "做法：",
        "1. 打開各樓層的「航廈對照圖_*.pdf」，圖上每個房間／房名都有編號（例如 3F-017）。",
        "2. 在「配對表」的「圖上編號」欄，填入該存置地點在圖上的編號；找不到對應的留空即可。",
        "3. 若圖上沒有範圍（開放大廳等），請在「備註」欄用格線描述，例如「D4–F6」，或寫「大廳北側」。",
        "4. 「我先前的推測」是我之前猜的位置，對的話可直接複製到「圖上編號」，不對的請覆寫。",
        "5. 同一個地點跨多個房間時，編號用逗號隔開（3F-017, 3F-018）。",
        "6. 若「資料推測樓層」是 ?，請在「樓層」欄寫 B1／淺挖區／1F／2F／3F。",
        "",
        "圖面方位：圖的左邊是北邊（每頁右上角有紅色北向箭頭）。",
        "格線：每格 20 m，欄用字母、列用數字（例如 D4）；PDF 是向量圖，可以無限放大，也可以搜尋房名。",
        "橘色菱形：圖上有房名文字、但自動偵測不到房間範圍。",
        "樓層名稱：「B1 停車場」＝檔名「航廈地下一樓」（圖面標題是地下室貳樓）；「淺挖區」＝檔名「淺漥區」。",
    ]:
        note.append([line])
    note.column_dimensions["A"].width = 110
    note["A1"].font = Font(bold=True, size=14)
    wb.save(OUT / "航廈配對表.xlsx")
    print("saved 航廈配對表.xlsx:", ws.max_row - 1, "locations")


if __name__ == "__main__":
    main()
